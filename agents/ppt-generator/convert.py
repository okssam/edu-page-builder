#!/usr/bin/env python3
"""
HTML 슬라이드를 PPTX로 변환하는 스크립트.

사용법:
    python3 convert.py <input.html> [output.pptx]

필요 패키지:
    pip install python-pptx beautifulsoup4
"""

import sys
import os
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("python-pptx가 설치되지 않았습니다.")
    print("설치: pip install python-pptx beautifulsoup4")
    sys.exit(1)

try:
    from bs4 import BeautifulSoup
except ImportError:
    print("beautifulsoup4가 설치되지 않았습니다.")
    print("설치: pip install beautifulsoup4")
    sys.exit(1)


# 스타일 가이드 색상
COLORS = {
    "primary": RGBColor(0x1E, 0x3A, 0x5F),
    "secondary": RGBColor(0x4A, 0x90, 0xD9),
    "accent": RGBColor(0xF3, 0x9C, 0x12),
    "text_primary": RGBColor(0x2C, 0x3E, 0x50),
    "text_secondary": RGBColor(0x7F, 0x8C, 0x8D),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}


def parse_html_slides(html_path: str) -> list[dict]:
    """HTML 파일에서 슬라이드 섹션들을 파싱합니다."""
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()

    soup = BeautifulSoup(html, "html.parser")
    slides = []

    for section in soup.find_all("section", class_="slide"):
        slide_data = {
            "number": section.get("data-slide-number", ""),
            "type": section.get("data-slide-type", "content"),
            "elements": [],
        }

        # 제목 추출
        for heading in section.find_all(["h1", "h2", "h3"]):
            slide_data["elements"].append({
                "type": "heading",
                "tag": heading.name,
                "text": heading.get_text(strip=True),
            })

        # 본문 텍스트 추출
        for p in section.find_all("p"):
            text = p.get_text(strip=True)
            if text:
                slide_data["elements"].append({
                    "type": "paragraph",
                    "text": text,
                })

        # 리스트 추출
        for ul in section.find_all(["ul", "ol"]):
            items = [li.get_text(strip=True) for li in ul.find_all("li")]
            slide_data["elements"].append({
                "type": "list",
                "items": items,
                "ordered": ul.name == "ol",
            })

        slides.append(slide_data)

    return slides


def create_pptx(slides: list[dict], output_path: str):
    """파싱된 슬라이드 데이터를 PPTX로 변환합니다."""
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 16:9
    prs.slide_height = Emu(6858000)

    blank_layout = prs.slide_layouts[6]  # 빈 레이아웃

    for slide_data in slides:
        slide = prs.slides.add_slide(blank_layout)
        slide_type = slide_data["type"]

        if slide_type == "cover":
            _build_cover_slide(slide, slide_data)
        elif slide_type == "objectives":
            _build_content_slide(slide, slide_data, is_objectives=True)
        else:
            _build_content_slide(slide, slide_data)

        # 슬라이드 번호 추가 (표지 제외)
        if slide_type != "cover" and slide_data["number"]:
            _add_slide_number(slide, slide_data["number"])

    prs.save(output_path)
    print(f"PPTX 생성 완료: {output_path}")


def _build_cover_slide(slide, data: dict):
    """표지 슬라이드를 구성합니다."""
    # 배경색
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLORS["primary"]

    y_offset = Inches(2.0)
    for elem in data["elements"]:
        if elem["type"] == "heading":
            txBox = slide.shapes.add_textbox(
                Inches(1), y_offset, Inches(11), Inches(1.2)
            )
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = elem["text"]
            p.font.size = Pt(44) if elem["tag"] == "h1" else Pt(24)
            p.font.bold = elem["tag"] == "h1"
            p.font.color.rgb = COLORS["white"]
            p.alignment = PP_ALIGN.CENTER
            y_offset += Inches(1.0)
        elif elem["type"] == "paragraph":
            txBox = slide.shapes.add_textbox(
                Inches(1), y_offset, Inches(11), Inches(0.6)
            )
            tf = txBox.text_frame
            p = tf.paragraphs[0]
            p.text = elem["text"]
            p.font.size = Pt(18)
            p.font.color.rgb = COLORS["white"]
            p.alignment = PP_ALIGN.CENTER
            y_offset += Inches(0.6)


def _build_content_slide(slide, data: dict, is_objectives: bool = False):
    """내용 슬라이드를 구성합니다."""
    y_offset = Inches(0.5)

    for elem in data["elements"]:
        if elem["type"] == "heading":
            txBox = slide.shapes.add_textbox(
                Inches(0.8), y_offset, Inches(11), Inches(0.8)
            )
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = elem["text"]
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = COLORS["primary"]
            y_offset += Inches(0.9)

        elif elem["type"] == "paragraph":
            txBox = slide.shapes.add_textbox(
                Inches(0.8), y_offset, Inches(11), Inches(0.6)
            )
            tf = txBox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = elem["text"]
            p.font.size = Pt(18)
            p.font.color.rgb = COLORS["text_primary"]
            y_offset += Inches(0.5)

        elif elem["type"] == "list":
            txBox = slide.shapes.add_textbox(
                Inches(1.0), y_offset, Inches(10.5), Inches(3.5)
            )
            tf = txBox.text_frame
            tf.word_wrap = True

            for i, item in enumerate(elem["items"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                prefix = f"{i+1}. " if elem["ordered"] else "  \u2022  "
                p.text = prefix + item
                p.font.size = Pt(18)
                p.font.color.rgb = COLORS["text_primary"]
                p.space_after = Pt(8)

            y_offset += Inches(0.4 * len(elem["items"]))


def _add_slide_number(slide, number: str):
    """슬라이드 우하단에 번호를 추가합니다."""
    txBox = slide.shapes.add_textbox(
        Inches(11.5), Inches(6.8), Inches(1), Inches(0.3)
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = str(number)
    p.font.size = Pt(11)
    p.font.color.rgb = COLORS["text_secondary"]
    p.alignment = PP_ALIGN.RIGHT


def main():
    if len(sys.argv) < 2:
        print("사용법: python3 convert.py <input.html> [output.pptx]")
        sys.exit(1)

    input_path = sys.argv[1]
    if not os.path.exists(input_path):
        print(f"파일을 찾을 수 없습니다: {input_path}")
        sys.exit(1)

    if len(sys.argv) >= 3:
        output_path = sys.argv[2]
    else:
        output_path = str(Path(input_path).with_suffix(".pptx"))

    slides = parse_html_slides(input_path)
    if not slides:
        print("슬라이드를 찾을 수 없습니다. HTML에 <section class='slide'> 태그가 있는지 확인하세요.")
        sys.exit(1)

    print(f"총 {len(slides)}개 슬라이드를 변환합니다...")
    create_pptx(slides, output_path)


if __name__ == "__main__":
    main()
